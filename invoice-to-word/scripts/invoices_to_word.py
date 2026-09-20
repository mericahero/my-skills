from pathlib import Path
import argparse
import re
import shutil

import pymupdf
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm


def parse_args():
    parser = argparse.ArgumentParser(description="Convert invoice PDFs to a Word document.")
    parser.add_argument("--source", default=".", help="Directory containing invoice PDFs.")
    parser.add_argument("--output", default=None, help="Output .docx path.")
    parser.add_argument("--include", action="append", default=None, help="Glob pattern(s) to select PDFs, relative to --source.")
    parser.add_argument("--dpi", type=int, default=220, help="Rendering DPI for invoice images.")
    parser.add_argument("--keep-images", action="store_true", help="Keep rendered PNG files after generating the Word document.")
    return parser.parse_args()


def pdf_files(source: Path, include_patterns):
    if include_patterns:
        files = []
        for pattern in include_patterns:
            files.extend(source.glob(pattern))
        files = sorted(set(files))
    else:
        files = sorted(source.glob("*.pdf"))
    return [f for f in files if f.is_file()]


def invoice_date(path: Path):
    with pymupdf.open(path) as doc:
        text = doc[0].get_text()
    match = re.search(r"(\d{4})年\s*(\d{1,2})月\s*(\d{1,2})日", text)
    if match:
        year, month, day = map(int, match.groups())
        return (year, month, day, path.name)
    return (9999, 12, 31, path.name)


def render_pdf(pdf_path: Path, png_path: Path, dpi: int):
    with pymupdf.open(pdf_path) as doc:
        page = doc[0]
        matrix = pymupdf.Matrix(dpi / 72, dpi / 72)
        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
        pixmap.save(png_path)


def image_dimensions(png_path: Path):
    with Image.open(png_path) as image:
        return image.size


def build_docx(image_paths, output_path: Path):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    for index, image_path in enumerate(image_paths):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Cm(0)
        paragraph.paragraph_format.space_after = Cm(0)
        paragraph.paragraph_format.line_spacing = 1
        if index > 0 and index % 2 == 0:
            paragraph.paragraph_format.page_break_before = True

        width, height = image_dimensions(image_path)
        aspect = width / height
        if aspect >= 1:
            image_width = Cm(17.5)
        else:
            image_width = Cm(12 * aspect)

        run = paragraph.add_run()
        run.add_picture(str(image_path), width=image_width)

        if index % 2 == 0:
            paragraph.paragraph_format.space_after = Cm(1.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main():
    args = parse_args()
    source = Path(args.source).expanduser().resolve()
    if not source.is_dir():
        raise SystemExit(f"Source directory does not exist: {source}")

    output = Path(args.output).expanduser().resolve() if args.output else source / "发票.docx"
    if output.suffix.lower() != ".docx":
        raise SystemExit("Output path must end with .docx")

    files = pdf_files(source, args.include)
    if not files:
        raise SystemExit(f"No PDF files found in: {source}")

    files = sorted(files, key=invoice_date)

    image_dir = source / "_invoice_images"
    image_dir.mkdir(parents=True, exist_ok=True)
    image_paths = []
    for index, pdf_path in enumerate(files, start=1):
        png_path = image_dir / f"invoice_{index:04d}.png"
        render_pdf(pdf_path, png_path, args.dpi)
        image_paths.append(png_path)

    build_docx(image_paths, output)

    if not args.keep_images:
        shutil.rmtree(image_dir, ignore_errors=True)

    print(f"Generated {output} with {len(files)} invoices.")


if __name__ == "__main__":
    main()
